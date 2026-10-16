"""Excel export into an existing workbook: only one sheet is replaced, the rest stays as it is."""

import base64
import datetime as dt
import io
import posixpath
import re
import sys
import zipfile

import openpyxl
import polars as pl
import pytest
import xlsxwriter
from lxml import etree

from tafl.api import ActionError
from tafl.api.testing import StubContext
from tafl.api.xlsx import _follow

DATA = pl.DataFrame(
    {
        "Customer": ["Huber", "Müller & <Co>", "=not a formula"],
        "Day": [dt.date(2025, 1, 2), dt.date(2025, 1, 3), None],
        "Net": [10.5, float("nan"), 3.25],
        "Orders": [1, 2, 3],
    }
)
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELATIONSHIP = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


@pytest.fixture
def report(tmp_path):
    """A report workbook: a sheet with formulas, a chart, a picture, a shape and a comment; the
    data sheet with the Excel table Sales and named ranges on it; a notes sheet. Plus a pivot
    table's data source on the data sheet and Excel's list of formula cells."""
    path = tmp_path / "report.xlsx"
    book = xlsxwriter.Workbook(path)
    page, data, notes = book.add_worksheet("Report"), book.add_worksheet("Data"), book.add_worksheet("Notes")
    page.write("A1", "Sales report", book.add_format({"bold": True, "font_size": 16, "bg_color": "yellow"}))
    page.write_formula("A2", "=SUM(Sales[Net])", None, 10)
    page.write_comment("A1", "made by hand")
    data.add_table(
        "A1:B5",
        {
            "name": "Sales",
            "columns": [{"header": "Customer"}, {"header": "Net"}],
            "data": [["old A", 1], ["old B", 2], ["old C", 3], ["old D", 4]],
        },
    )
    data.set_tab_color("red")
    chart = book.add_chart({"type": "column"})
    chart.add_series({"categories": "=Data!$A$2:$A$5", "values": "=Data!$B$2:$B$5", "name": "=Data!$B$1"})
    page.insert_chart("D2", chart)
    page.insert_image("D20", "dot.png", {"image_data": io.BytesIO(PNG)})
    page.insert_textbox("H2", "A note in a shape")
    book.define_name("SalesData", "=Data!$A$1:$B$5")
    book.define_name("Nets", "=Data!$B$2:$B$5")
    notes.write("A1", "keep me")
    book.close()
    _add_parts(
        path,
        {
            "xl/pivotCache/pivotCacheDefinition1.xml": (
                f'<pivotCacheDefinition xmlns="{MAIN}"><cacheSource type="worksheet">'
                '<worksheetSource ref="A1:B5" sheet="Data"/></cacheSource><cacheFields count="0"/></pivotCacheDefinition>',
                "pivotCacheDefinition",
            ),
            "xl/calcChain.xml": (f'<calcChain xmlns="{MAIN}"><c r="A2" i="1"/></calcChain>', "calcChain"),
        },
    )
    return path


def test_only_the_data_sheet_is_replaced(act, report):
    before = _parts(report)
    params = {"file": str(report), "sheet": "data", "existing": "sheet", "table_style": "Table Style Medium 2"}
    act("excel.export", DATA, params)
    after = _parts(report)
    _check_package(report)
    assert set(before) - set(after) == {"xl/calcChain.xml"} and set(after) <= set(before)
    changed = {name for name in after if after[name] != before[name]}
    assert changed == {
        "[Content_Types].xml",
        "xl/_rels/workbook.xml.rels",  # without the list of formula cells
        "xl/workbook.xml",  # named ranges, calculation when opened
        "xl/styles.xml",
        "xl/worksheets/sheet2.xml",
        "xl/worksheets/_rels/sheet2.xml.rels",
        "xl/tables/table1.xml",
        "xl/charts/chart1.xml",
        "xl/pivotCache/pivotCacheDefinition1.xml",
    }  # the other sheets, the drawing with chart, picture and shape, the comment: all as they were

    book = openpyxl.load_workbook(report)
    assert book.sheetnames == ["Report", "Data", "Notes"]  # in its place, with its name's spelling
    data = book["Data"]
    assert data.sheet_properties.tabColor.rgb == "FFFF0000"
    assert [c.value for c in data[1]] == ["Customer", "Day", "Net", "Orders"] and data["A1"].font.b
    assert [c.value for c in data[3]] == ["Müller & <Co>", dt.datetime(2025, 1, 3), None, 2]  # NaN: empty
    assert data["A4"].value == "=not a formula" and data["A4"].data_type == "s"
    assert data["B2"].number_format == "dd.mm.yyyy" and data["C2"].number_format == "#,##0.00"
    assert data.freeze_panes == "A2"
    (table,) = data.tables.values()  # keeps its name: =SUM(Sales[Net]) still works
    assert (table.displayName, table.ref, table.tableStyleInfo.name) == ("Sales", "A1:D4", "TableStyleMedium2")
    page = book["Report"]
    assert page["A1"].font.b and page["A1"].fill.fgColor.rgb == "FFFFFF00" and page["A2"].value == "=SUM(Sales[Net])"
    assert book.calculation.fullCalcOnLoad
    # ranges over the old rows 2 to 5 now end at the new last row 4
    assert book.defined_names["SalesData"].attr_text == "Data!$A$1:$B$4"
    assert book.defined_names["Nets"].attr_text == "Data!$B$2:$B$4"
    chart = after["xl/charts/chart1.xml"].decode()
    assert re.findall(r"<c:f>(.*?)</c:f>", chart) == ["Data!$B$1", "Data!$A$2:$A$4", "Data!$B$2:$B$4"]
    pivot = etree.fromstring(after["xl/pivotCache/pivotCacheDefinition1.xml"])
    assert pivot.get("refreshOnLoad") == "1" and pivot.find(f".//{{{MAIN}}}worksheetSource").get("ref") == "A1:B4"

    back = act("excel.import", None, {"file": str(report), "sheet": "Data"})
    assert back["Customer"].to_list() == DATA["Customer"].to_list() and back["Day"][0] == dt.date(2025, 1, 2)


def test_a_new_sheet_is_added_at_the_end(act, report):
    act("excel.export", DATA, {"file": str(report), "sheet": "Sales", "existing": "sheet"})
    _check_package(report)
    book = openpyxl.load_workbook(report)
    assert book.sheetnames == ["Report", "Data", "Notes", "Sales"]
    assert list(book["Data"].tables) == ["Sales"] and list(book["Sales"].tables) == ["Sales_2"]  # names differ
    assert not book["Sales"].sheet_view.tabSelected  # otherwise Excel would group it with the selected sheet
    assert book.defined_names["Nets"].attr_text == "Data!$B$2:$B$5"  # the data sheet didn't change


def test_formats_are_added_only_once(act, report):
    params = {"file": str(report), "sheet": "Data", "existing": "sheet"}
    act("excel.export", DATA, params)
    styles = _parts(report)["xl/styles.xml"]
    act("excel.export", DATA, params)
    assert _parts(report)["xl/styles.xml"] == styles


def test_workbook_with_macros_keeps_them(act, tmp_path):
    path = tmp_path / "macros.xlsm"
    project = tmp_path / "vbaProject.bin"
    project.write_bytes(b"not a real macro project, it only has to stay")
    book = xlsxwriter.Workbook(path)
    book.add_worksheet("Start").write("A1", "Start")
    data = book.add_worksheet("Data")
    data.set_vba_name("DataSheet")
    data.hide()
    book.add_vba_project(str(project))
    book.close()

    act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"})
    _check_package(path)
    with zipfile.ZipFile(path) as archive:
        assert archive.read("xl/vbaProject.bin") == project.read_bytes()
        assert "macroEnabled" in archive.read("[Content_Types].xml").decode()
    book = openpyxl.load_workbook(path)
    data = book["Data"]
    assert data.sheet_state == "hidden" and data.sheet_properties.codeName == "DataSheet"
    assert data["A2"].value == "Huber"

    with pytest.raises(ActionError, match="macros"):  # Tafl can't write a whole workbook with macros
        act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "file"})
    with pytest.raises(ActionError, match="macros"):
        act("excel.export", DATA, {"file": str(tmp_path / "new.xlsm"), "existing": "sheet"})


def test_dates_in_a_workbook_with_1904_dates(act, tmp_path):
    path = tmp_path / "mac.xlsx"
    book = xlsxwriter.Workbook(path, {"date_1904": True})
    book.add_worksheet("Data")
    book.close()
    act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"})
    data = openpyxl.load_workbook(path)["Data"]
    assert data["B2"].value == dt.datetime(2025, 1, 2)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows locks open files")
def test_an_open_workbook_is_left_alone(act, report):
    before = report.read_bytes()
    with open(report, "rb"), pytest.raises(ActionError, match="open in Excel"):
        act("excel.export", DATA, {"file": str(report), "sheet": "Data", "existing": "sheet"})
    assert report.read_bytes() == before
    assert [p.name for p in report.parent.iterdir()] == ["report.xlsx"]  # no temporary file left


def test_not_a_sheet_with_cells_or_not_a_workbook(act, tmp_path):
    path = tmp_path / "charts.xlsx"
    book = xlsxwriter.Workbook(path)
    book.add_worksheet("Data").write_column("A1", [1, 2])
    chart = book.add_chart({"type": "line"})
    chart.add_series({"values": "=Data!$A$1:$A$2"})
    book.add_chartsheet("Chart").set_chart(chart)
    book.close()
    with pytest.raises(ActionError, match="chart sheet"):
        act("excel.export", DATA, {"file": str(path), "sheet": "Chart", "existing": "sheet"})
    text = tmp_path / "text.xlsx"
    text.write_text("no workbook")
    with pytest.raises(ActionError, match="not an Excel workbook"):
        act("excel.export", DATA, {"file": str(text), "sheet": "Data", "existing": "sheet"})


def test_without_a_file_a_new_workbook_is_written(act, tmp_path):
    path = tmp_path / "new.xlsx"
    act("excel.export", DATA, {"file": str(path), "sheet": "2025", "existing": "sheet"})
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["2025"] and list(book["2025"].tables) == ["Table_2025"]


def test_old_steps_replace_the_whole_file(registry):
    assert registry.actions["excel.export"].upgrade({"file": "x.xlsx"}, 1)["existing"] == "file"


def test_ranges_that_follow_the_data():
    formula = "SUM('My data'!$B$2:$B$5,Other!B2:B5,'My data'!B3:B5,[1]Data!B2:B5,'My data'!A1:D5)"
    assert _follow(formula, "my data", 5, 9) == (
        "SUM('My data'!$B$2:$B$9,Other!B2:B5,'My data'!B3:B5,[1]Data!B2:B5,'My data'!A1:D9)"
    )


def test_a_template_with_only_the_header_row(act, tmp_path):
    """Ranges down to row 2, the empty row left for the data, follow the data; row 1 alone stays."""
    path = tmp_path / "template.xlsx"
    book = xlsxwriter.Workbook(path)
    book.add_worksheet("Data").write_row("A1", ["Customer", "Day", "Net", "Orders"])
    chart = book.add_chart({"type": "column"})
    chart.add_series({"values": "=Data!$C$2:$C$2"})
    book.add_worksheet("Report").insert_chart("A1", chart)
    book.define_name("Nets", "=Data!$C$2:$C$2")
    book.define_name("Header", "=Data!$A$1:$D$1")
    book.close()
    act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"})
    names = {n: v.attr_text for n, v in openpyxl.load_workbook(path).defined_names.items()}
    assert names == {"Nets": "Data!$C$2:$C$4", "Header": "Data!$A$1:$D$1"}
    assert b"Data!$C$2:$C$4" in _parts(path)["xl/charts/chart1.xml"]


def test_what_the_replaced_sheet_had_of_its_own_is_named(act, tmp_path):
    path = tmp_path / "data.xlsx"
    book = xlsxwriter.Workbook(path)
    data = book.add_worksheet("Data")
    data.write_column("A1", ["Net", 1, 2])
    chart = book.add_chart({"type": "line"})
    chart.add_series({"values": "=Data!$A$2:$A$3"})
    data.insert_chart("C2", chart)
    data.insert_image("C20", "dot.png", {"image_data": io.BytesIO(PNG)})
    data.insert_textbox("H2", "a shape")
    data.write_comment("A1", "a note")
    data.conditional_format("A2:A3", {"type": "cell", "criteria": ">", "value": 1, "format": book.add_format({"bold": True})})
    data.data_validation("B2:B3", {"validate": "integer", "criteria": ">", "value": 0})
    book.close()
    ctx = StubContext()
    act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"}, ctx=ctx)
    assert ctx.messages[-2] == (
        "warning",
        "The old sheet 'Data' had its own charts, pictures, shapes, comments, conditional formats and data validation, "
        "and they were removed with it (the other sheets keep theirs). To keep such things, put them on another sheet.",
    )
    _check_package(path)
    assert openpyxl.load_workbook(path)["Data"]["A2"].value == "Huber"


@pytest.mark.parametrize("kind", ["pivotTable", "slicer", "timeline"])
def test_a_sheet_with_its_own_pivot_tables_or_slicers_is_not_replaced(act, tmp_path, kind):
    """Excel connects them with caches elsewhere in the workbook: Tafl leaves such a sheet alone."""
    path = tmp_path / "pivot.xlsx"
    book = xlsxwriter.Workbook(path)
    book.add_worksheet("Data").write_column("A1", ["Net", 1, 2])
    book.close()
    contents = _parts(path)
    contents[f"xl/{kind}s/{kind}1.xml"] = b"<x/>"
    folder = {"slicer": "http://schemas.microsoft.com/office/2007/relationships/",
              "timeline": "http://schemas.microsoft.com/office/2011/relationships/"}.get(kind, RELATIONSHIP)  # fmt: skip
    contents["xl/worksheets/_rels/sheet1.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{folder}{kind}" Target="../{kind}s/{kind}1.xml"/></Relationships>'
    ).encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    before = path.read_bytes()
    word = {"pivotTable": "pivot tables", "slicer": "slicers", "timeline": "timelines"}[kind]
    with pytest.raises(ActionError, match=f"has {word} of its own.*Move them to another sheet"):
        act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"})
    assert path.read_bytes() == before


def test_lost_objects_that_are_easy_to_mistake(act, tmp_path):
    """A new kind of chart with its stand-in text box, an old form button, a header picture and
    an empty list of validation rules; a drawing that is referred to but missing."""
    path = tmp_path / "data.xlsx"
    book = xlsxwriter.Workbook(path)
    book.add_worksheet("Data").write_column("A1", ["Net", 1])
    book.close()
    contents = _parts(path)
    drawing = (
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><xdr:twoCellAnchor>'
        '<mc:AlternateContent><mc:Choice Requires="cx1"><xdr:graphicFrame/></mc:Choice>'
        "<mc:Fallback><xdr:sp><xdr:txBody>This chart isn't available</xdr:txBody></xdr:sp></mc:Fallback>"
        "</mc:AlternateContent></xdr:twoCellAnchor></xdr:wsDr>"
    )
    contents["xl/drawings/drawing1.xml"] = drawing.encode()
    contents["xl/drawings/_rels/drawing1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2014/relationships/chartEx" Target="../charts/chartEx1.xml"/>'
        b"</Relationships>"
    )
    contents["xl/charts/chartEx1.xml"] = b"<x/>"
    contents["xl/drawings/vmlDrawing1.vml"] = b'<xml><v:shape><x:ClientData ObjectType="Button"/></v:shape></xml>'
    relations = (
        f'<Relationship Id="rId1" Type="{RELATIONSHIP}drawing" Target="../drawings/drawing1.xml"/>'
        f'<Relationship Id="rId2" Type="{RELATIONSHIP}vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
        f'<Relationship Id="rId3" Type="{RELATIONSHIP}drawing" Target="../drawings/missing.xml"/>'
    )
    contents["xl/worksheets/_rels/sheet1.xml.rels"] = (
        f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{relations}</Relationships>'
    ).encode()
    sheet = contents["xl/worksheets/sheet1.xml"].decode()
    contents["xl/worksheets/sheet1.xml"] = sheet.replace(
        "</worksheet>", '<dataValidations count="0"/><legacyDrawingHF r:id="rId9"/></worksheet>'
    ).encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    ctx = StubContext()
    act("excel.export", DATA, {"file": str(path), "sheet": "Data", "existing": "sheet"}, ctx=ctx)
    warning = next(text for level, text in ctx.messages if level == "warning")
    assert warning.startswith("The old sheet 'Data' had its own charts, form controls and header and footer pictures,")


def test_a_tag_cut_in_two_by_a_piece_is_found():
    from tafl.api.xlsx import _SHEET_OWN, _search

    for offset in (-20, -5, 0, 5):  # around the end of the first 1 MiB piece
        text = b"x" * ((1 << 20) + offset) + b"<conditionalFormatting sqref='A1'>"
        assert _search(io.BytesIO(text), _SHEET_OWN) == [True, False, False], offset


def _parts(path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _add_parts(path, parts: dict[str, tuple[str, str]]) -> None:
    """Adds parts to the workbook, each with a relationship from the workbook part."""
    contents = _parts(path)
    rels, types = contents["xl/_rels/workbook.xml.rels"].decode(), contents["[Content_Types].xml"].decode()
    for number, (name, (xml, kind)) in enumerate(parts.items(), start=90):
        contents[name] = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n{xml}'.encode()
        target = posixpath.relpath(name, "xl")
        relation = f'<Relationship Id="rId{number}" Type="{RELATIONSHIP}{kind}" Target="{target}"/>'
        rels = rels.replace("</Relationships>", relation + "</Relationships>")
        content_type = f"application/vnd.openxmlformats-officedocument.spreadsheetml.{kind}+xml"
        types = types.replace("</Types>", f'<Override PartName="/{name}" ContentType="{content_type}"/></Types>')
    contents["xl/_rels/workbook.xml.rels"], contents["[Content_Types].xml"] = rels.encode(), types.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in contents.items():
            archive.writestr(name, content)


def _check_package(path) -> None:
    """Every relationship leads to a part, every part has a content type and every content type
    belongs to a part."""
    parts = _parts(path)
    types = etree.fromstring(parts["[Content_Types].xml"])
    defaults = {d.get("Extension").lower() for d in types if d.tag.endswith("Default")}
    overrides = {o.get("PartName").lstrip("/").lower() for o in types if o.tag.endswith("Override")}
    names = {name.lower() for name in parts}
    assert overrides <= names, overrides - names
    for name in parts:
        if name != "[Content_Types].xml":
            assert name.lower() in overrides or name.rsplit(".", 1)[-1].lower() in defaults, name
        if name.endswith(".rels"):
            folder = posixpath.dirname(posixpath.dirname(name))
            source = posixpath.join(folder, posixpath.basename(name)[: -len(".rels")])
            assert source == "" or source.lower() in names, name
            for relation in etree.fromstring(parts[name]):
                if relation.get("TargetMode") != "External":
                    target = relation.get("Target")
                    part = target[1:] if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
                    assert part.lower() in names, (name, target)
