import datetime as dt
import decimal
import logging
import sys
from pathlib import Path

import openpyxl
import polars as pl
import pytest
import xlsxwriter

from tafl.api import ActionError
from tafl.api.testing import StubContext


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "report.xlsx"
    wb = xlsxwriter.Workbook(path)
    date_format = wb.add_format({"num_format": "dd.mm.yyyy"})
    ws = wb.add_worksheet("Daten")
    ws.write("A1", "Monthly report")
    ws.write_row("A3", ["Kunde", "Menge", "Preis", "Datum", None, "Notiz"])
    ws.write_row("A4", ["Anna", 3, 9.99])
    ws.write_datetime("D4", dt.datetime(2025, 1, 31), date_format)
    ws.write("F4", "urgent")
    ws.write_row("A5", ["Ben", 5, 1.5])
    ws.write_datetime("D5", dt.datetime(2025, 2, 1), date_format)
    ws.write("F5", 12)
    wb.add_worksheet("Other").write_row("A1", ["x"])
    wb.close()
    return path


def test_import_sheet_with_header_further_down(act, workbook):
    df = act("excel.import", None, {"file": str(workbook), "sheet": "daten", "header_row": 3})
    assert df.columns == ["Kunde", "Menge", "Preis", "Datum", "Notiz"]  # empty unnamed column dropped
    assert df.schema == {"Kunde": pl.String, "Menge": pl.Int64, "Preis": pl.Float64, "Datum": pl.Date, "Notiz": pl.String}
    assert df["Datum"].to_list() == [dt.date(2025, 1, 31), dt.date(2025, 2, 1)]
    assert df["Notiz"].to_list() == ["urgent", "12"]


def test_unknown_sheet_lists_the_sheets(act, workbook):
    with pytest.raises(ActionError, match="Sheets: Daten, Other"):
        act("excel.import", None, {"file": str(workbook), "sheet": "Nope"})


def test_export_with_formats(act, tmp_path):
    df = pl.DataFrame(
        {
            "Name": ["Anna", "Ben"],
            "Qty": [3, 5],
            "Price": [1234.5, 0.25],
            "Rate": [0.19, 0.07],
            "Day": [dt.date(2025, 12, 31), None],
        }
    )
    path = tmp_path / "out.xlsx"
    act(
        "excel.export",
        df,
        {"file": str(path), "sheet": "Orders", "column_formats": [{"column": "Rate", "format": "0%"}]},
    )
    ws = openpyxl.load_workbook(path)["Orders"]
    assert [c.value for c in ws[1]] == ["Name", "Qty", "Price", "Rate", "Day"]
    assert ws["A1"].font.b
    assert ws.freeze_panes == "A2"
    assert ws["B2"].number_format == "0"
    assert ws["C2"].number_format == "#,##0.00"
    assert ws["D2"].number_format == "0%"
    assert ws["E2"].number_format == "dd.mm.yyyy"
    assert ws["E2"].value == dt.datetime(2025, 12, 31)


def test_export_checks_file_and_sheet_names(act, tmp_path):
    df = pl.DataFrame({"a": [1]})
    with pytest.raises(ActionError, match=r"\.xlsx"):
        act("excel.export", df, {"file": str(tmp_path / "out.xls")})
    with pytest.raises(ActionError, match="31 characters"):
        act("excel.export", df, {"file": str(tmp_path / "out.xlsx"), "sheet": "a/b"})


@pytest.fixture
def edges(tmp_path):
    """A workbook as other programs write it: merged cells, formulas without saved results,
    cells with only a time."""
    path = tmp_path / "edges.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet.append(["Region", "Name", "Summe", "Zeit"])
    for row in (["Nord", "a", "=1+1", dt.time(9, 30)], [None, "b", "=1/0", dt.time(10, 0)], ["Süd", "c", "=2+2", dt.time(11, 15)]):
        sheet.append(row)
    sheet.merge_cells("A2:A3")
    book.save(path)
    return path


def test_import_cells_with_only_a_time(act, edges, caplog):
    with caplog.at_level(logging.WARNING):
        df = act("excel.import", None, {"file": str(edges)})
    assert df.schema["Zeit"] == pl.Time
    assert df["Zeit"].to_list() == [dt.time(9, 30), dt.time(10, 0), dt.time(11, 15)]
    assert df["Region"].to_list() == ["Nord", None, "Süd"]  # merged cells: the value is in the first cell only
    assert df["Summe"].to_list() == [None, None, None]  # formulas without saved results
    assert not [r for r in caplog.records if r.name.startswith("fastexcel")]  # its note on the console about them


def test_a_header_row_below_the_data(act, edges):
    ctx = StubContext()
    df = act("excel.import", None, {"file": str(edges), "header_row": 100}, ctx)
    assert df.shape == (0, 0)
    assert ctx.warnings == ["Sheet 'S' of edges.xlsx has nothing in row 100 (the row with the column names) or below"]


def test_export_binary_values_decimals_and_long_texts(act, tmp_path):
    df = pl.DataFrame(
        {
            "data": [b"\x01\x02\xff", None],
            "price": pl.Series([decimal.Decimal("1.50"), decimal.Decimal("2.25")], dtype=pl.Decimal(10, 2)),
            "text": ["x" * 40_000, "short"],
        }
    )
    ctx = StubContext()
    path = tmp_path / "out.xlsx"
    act("excel.export", df, {"file": str(path)}, ctx)
    ws = openpyxl.load_workbook(path).active
    assert [ws["A2"].value, ws["A3"].value] == ["0x0102FF", None]
    assert ws["B2"].value == 1.5 and ws["B2"].number_format == "#,##0.00"
    assert len(ws["C2"].value) == 32_767
    assert ctx.warnings == ["1 text was cut to 32,767 characters, the most an Excel cell holds"]


def test_export_more_rows_than_a_sheet_holds(act, tmp_path):
    with pytest.raises(ActionError, match=r"The table has 1,048,576 rows, but an Excel sheet holds at most 1,048,575"):
        act("excel.export", pl.DataFrame({"a": range(1_048_576)}), {"file": str(tmp_path / "big.xlsx")})
    assert not list(tmp_path.iterdir())


def test_the_help_fits_both_choices_for_an_existing_file(registry):
    existing = next(p for p in registry.actions["excel.export"].params if p.key == "existing")
    assert existing.help.startswith("Replace the whole file: a new workbook with only this sheet. Replace only this sheet:")


def test_dont_export_an_empty_table(act, registry, tmp_path):
    path, empty = tmp_path / "out.xlsx", pl.DataFrame(schema={"a": pl.Int64})
    act("excel.export", pl.DataFrame({"a": [1]}), {"file": str(path)})
    before = path.read_bytes()
    for existing in ("file", "sheet"):
        ctx = StubContext()
        act("excel.export", empty, {"file": str(path), "existing": existing, "skip_empty": True}, ctx)
        assert path.read_bytes() == before and ctx.messages == [("info", "Not exported: the table is empty")]
    act("excel.export", empty, {"file": str(tmp_path / "new.xlsx"), "skip_empty": True})
    assert sorted(p.name for p in tmp_path.iterdir()) == ["out.xlsx"]
    act("excel.export", empty, {"file": str(path)})  # off: as before, the header only
    assert [[c.value for c in row] for row in openpyxl.load_workbook(path).active] == [["a"]]
    export = registry.actions["excel.export"]
    assert export.version == 3 and export.upgrade({"file": "a.xlsx"}, 2) == {"file": "a.xlsx", "skip_empty": False}
    assert export.upgrade({"file": "a.xlsx"}, 1) == {"file": "a.xlsx", "existing": "file", "skip_empty": False}


def test_a_table_without_columns_gives_an_empty_sheet(act, tmp_path):
    """As the CSV export gives an empty file, e.g. for a folder import without files and without
    its File column."""
    path = tmp_path / "out.xlsx"
    ctx = StubContext()
    act("excel.export", pl.DataFrame(), {"file": str(path), "sheet": "Data"}, ctx)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Data"] and book["Data"].max_column == 1 and book["Data"]["A1"].value is None
    assert ctx.messages == [("info", f"Wrote 0 rows to {path} (sheet 'Data')")]
    act("excel.export", pl.DataFrame({"a": [1]}), {"file": str(path), "sheet": "Other", "existing": "sheet"})
    act("excel.export", pl.DataFrame(), {"file": str(path), "sheet": "Other", "existing": "sheet"})
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Data", "Other"] and book["Other"]["A1"].value is None


def test_a_failed_export_leaves_the_old_file(act, tmp_path, monkeypatch):
    path = tmp_path / "out.xlsx"
    act("excel.export", pl.DataFrame({"a": [1]}), {"file": str(path)})
    before = path.read_bytes()

    def broken(self, workbook=None, **options):
        Path(workbook).write_bytes(b"PK")
        raise pl.exceptions.InvalidOperationError("broken")

    monkeypatch.setattr(pl.DataFrame, "write_excel", broken)
    with pytest.raises(pl.exceptions.InvalidOperationError):
        act("excel.export", pl.DataFrame({"a": [2]}), {"file": str(path)})
    assert path.read_bytes() == before
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]  # no temporary file left


@pytest.mark.skipif(sys.platform != "win32", reason="Windows locks open files")
def test_export_to_a_workbook_that_is_open(act, tmp_path):
    path = tmp_path / "out.xlsx"
    act("excel.export", pl.DataFrame({"a": [1]}), {"file": str(path)})
    before = path.read_bytes()
    with open(path, "rb"), pytest.raises(ActionError, match=r"out.xlsx is open in another program \(e.g. Excel\)"):
        act("excel.export", pl.DataFrame({"a": [2]}), {"file": str(path)})
    assert path.read_bytes() == before
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]
