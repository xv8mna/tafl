"""What a data sheet writes, read back through the real import actions."""

import contextlib
import json
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import polars as pl
import pytest
import xlsxwriter
from polars.testing import assert_frame_equal

from tafl.api import ActionError, Cancelled, files
from tafl.api.testing import StubContext
from tafl.core import Engine, Project
from tafl.core import sheet as module
from tafl.core.sheet import (
    FORMATS,
    Sheet,
    format_problem,
    mark,
    parse_text,
    read_file,
    source_format,
    stamp_of,
    write_file,
)

windows = pytest.mark.skipif(sys.platform != "win32", reason="Windows locks open files")
all_formats = pytest.mark.parametrize("suffix", list(FORMATS))

HARD_VALUES = [
    "00123", "1,50", "27.02.2026", "1E5", " leading", "trailing ", "   ", "=1+2", "'apos", '"quoted"', 'a"b', "a;b", "a\tb",
    "a\nb", "ÄÖÜ€ß", "\U0001f600", "_x0041_", "x\x01y", "TRUE", "1.0", "http://x.y/z", "#N/A", None, "-5", "+49 30", "0.1",
    "a,b", "\\N", "NA", "null", '""', '"', "a\n", "\nb", "\n", " ", "<a>&b", "--", "@x", "%d", "{x}", "é",
]  # fmt: skip


@contextlib.contextmanager
def held(path: Path, share_read: bool):
    """Holds a file open as another program does: Excel lets others read an open file, some
    programs don't."""
    import ctypes
    from ctypes import wintypes

    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CreateFileW.argtypes = [
        wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE,
    ]  # fmt: skip
    handle = kernel.CreateFileW(str(path), 0x80000000 | 0x40000000, 1 if share_read else 0, None, 3, 0, None)
    assert handle not in (None, wintypes.HANDLE(-1).value), ctypes.get_last_error()
    try:
        yield
    finally:
        kernel.CloseHandle(handle)


def text_frame(columns: dict[str, list]) -> pl.DataFrame:
    return pl.DataFrame(columns, schema={name: pl.String for name in columns})


def step_params(path: Path, suffix: str) -> dict:
    """The settings of the import step that a sheet adds."""
    sheet_format = FORMATS[suffix]
    return {**sheet_format.settings, "source": "file", "file": str(path), files.SHEET_KEY: mark(path)}


def read_back(registry, frame: pl.DataFrame, folder: Path, suffix: str) -> pl.DataFrame:
    path = folder / f"data{suffix}"
    write_file(frame, path, suffix)
    action = registry.actions[FORMATS[suffix].action]
    out = action(None, step_params(path, suffix), StubContext())
    read = read_file(path, FORMATS[suffix], action)
    assert_frame_equal(read, out)
    return out


@all_formats
def test_hard_values_read_back_exactly(registry, tmp_path, suffix):
    frame = text_frame({"Text": HARD_VALUES, "Same": HARD_VALUES[::-1], "Empty": [None] * len(HARD_VALUES)})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


XML_VALUES = [
    "<r>abc</r>", "<r>a&b</r>", "<r><t>x</t></r>", "<r>", "</r>", "<r></r>", "<r>line\nbreak</r>", "<r>  spaces  </r>",
    "<r>\x01</r>", "<r>_x0041_</r>", "<r>ÄÖÜ€</r>", "plain <r>x</r>", "<r>x</r> plain", "<R>x</R>",
]  # fmt: skip


@all_formats
def test_cells_that_look_like_rich_text_read_back_exactly(registry, tmp_path, suffix):
    frame = text_frame({"Text": XML_VALUES, "<r>a&b</r>": XML_VALUES[::-1]})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


def test_cells_that_look_like_rich_text_are_plain_text_in_the_workbook(tmp_path):
    frame = text_frame({"<r>a&b</r>": ["<r>a&b</r>", "<r>abc</r>", "<r><t>x</t></r>"]})
    path = tmp_path / "data.xlsx"
    write_file(frame, path, ".xlsx")
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    assert rows == [("<r>a&b</r>",), ("<r>a&b</r>",), ("<r>abc</r>",), ("<r><t>x</t></r>",)]


def test_the_longest_cell_that_looks_like_rich_text_reads_back_exactly(registry, tmp_path):
    long = "<r>" + "&" * (32_767 - 7) + "</r>"
    frame = text_frame({"Text": [long, "x"]})
    assert format_problem(frame, ".xlsx") is None
    assert_frame_equal(read_back(registry, frame, tmp_path, ".xlsx"), frame)


def test_a_column_name_like_a_generated_one_stays(registry, tmp_path):
    frame = text_frame({"a": ["1"], "__UNNAMED__5": ["2"], "__UNNAMED__0": ["3"]})
    assert_frame_equal(read_back(registry, frame, tmp_path, ".xlsx"), frame)


def test_a_header_cell_cleared_outside_gets_a_number(registry, tmp_path):
    path = tmp_path / "data.xlsx"
    write_file(text_frame({"a": ["1"], "b": ["2"]}), path, ".xlsx")
    book = openpyxl.load_workbook(path)
    book.active["B1"] = None
    book.save(path)
    action = registry.actions["excel.import"]
    assert read_file(path, FORMATS[".xlsx"], action).columns == ["a", "Column 2"]


@all_formats
def test_rows_without_a_value_stay(registry, tmp_path, suffix):
    frame = text_frame({"a": [None, "x", None, None, "y"], "b": [None, "1", None, None, None]})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_one_column_with_empty_rows(registry, tmp_path, suffix):
    frame = text_frame({"a": [None, "x", None, None, "1,50"]})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_a_table_with_no_rows(registry, tmp_path, suffix):
    frame = text_frame({"a": [], "b": []})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_a_column_without_a_value_stays_a_text_column(registry, tmp_path, suffix):
    sheet = Sheet(columns=2, rows=5)
    sheet.set_text(0, 0, " ")
    assert_frame_equal(read_back(registry, sheet.frame(), tmp_path, suffix), sheet.frame())


@all_formats
def test_names(registry, tmp_path, suffix):
    frame = text_frame(
        {
            "Column 1": ["1"],
            "id": ["2"],
            "id (2)": ["3"],
            "2025": ["4"],
            "a;b": ["5"],
            'q"r': ["6"],
            "a,b": ["7"],
            "Ü b": ["8"],
        }
    )
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_a_sheet_with_a_row_of_odd_names(registry, tmp_path, suffix):
    sheet = Sheet()
    sheet.load(parse_text("id\tid\t\t a \tb\\nc\n1\t2\t3\t4\t5\n"))
    sheet.set_header(True)
    assert sheet.names == ["id", "id (2)", "Column 3", "a", "b\\nc"]
    assert_frame_equal(read_back(registry, sheet.frame(), tmp_path, suffix), sheet.frame())


@all_formats
def test_what_was_pasted_reads_back_as_the_sheet_shows_it(registry, tmp_path, suffix):
    sheet = Sheet()
    sheet.load(
        parse_text(
            'Nr\tBetrag\tDatum\tText\r\n00123\t1,50\t27.02.2026\t"zwei\r\nZeilen"\r\n\t\t\t\r\n1E5\t 7 \t\t"""x"""\r\n'
        )
    )
    assert sheet.header
    out = read_back(registry, sheet.frame(), tmp_path, suffix)
    assert out.rows() == [
        ("00123", "1,50", "27.02.2026", "zwei\nZeilen"),
        (None, None, None, None),
        ("1E5", " 7 ", None, '"x"'),
    ]


@all_formats
def test_spare_rows_and_columns_are_left_out_and_empty_columns_with_a_name_stay(registry, tmp_path, suffix):
    sheet = Sheet()
    sheet.set_text(1, 0, "x")
    frame = sheet.frame()
    assert frame.rows() == [(None,), ("x",)]
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)
    sheet.rename_column(2, "Amount")
    frame = sheet.frame()
    assert frame.rows() == [(None, None, None), ("x", None, None)]
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_a_file_opens_in_a_sheet_as_it_was_written(registry, tmp_path, suffix):
    frame = text_frame({"Text": HARD_VALUES, "Same": HARD_VALUES[::-1]})
    out = read_back(registry, frame, tmp_path, suffix)
    sheet = Sheet.from_frame(out)
    assert not sheet.changed
    assert_frame_equal(sheet.frame(), frame)
    sheet.set_text(0, 0, "new")
    frame = sheet.frame()
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


@all_formats
def test_a_cell_with_a_carriage_return_in_a_file_stays(registry, tmp_path, suffix):
    """Such a cell can't be pasted, but a file can hold it."""
    frame = text_frame({"a": ["x\ry", "x\r\ny", "z"]})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


def test_a_file_with_numbers_from_elsewhere_opens_as_text(registry, tmp_path):
    path = tmp_path / "a.xlsx"
    pl.DataFrame({"n": [1, 2], "x": [1.5, None], "t": ["a", None]}).write_excel(path)
    out = read_file(path, FORMATS[".xlsx"], registry.actions["excel.import"])
    assert Sheet.from_frame(out).view().rows() == [("1", "1.5", "a"), ("2", None, None)]


def test_read_file_uses_its_own_settings_and_says_what_is_wrong(registry, tmp_path):
    path = tmp_path / "a.csv"
    write_file(text_frame({"a": ["1,50"]}), path, ".csv")
    assert read_file(path, FORMATS[".csv"], registry.actions["csv.import"]).rows() == [("1,50",)]
    path.unlink()
    with pytest.raises(ActionError, match=r"File not found"):
        read_file(path, FORMATS[".csv"], registry.actions["csv.import"])


def test_read_file_passes_on_what_the_import_left_out(registry, tmp_path):
    path = tmp_path / "a.csv"
    write_file(text_frame({"a": ["1"], "b": ["x"]}), path, ".csv")
    path.write_text(path.read_text(encoding="utf-8-sig") + "2;y;EXTRA\r\n", encoding="utf-8-sig")
    warnings = []
    out = read_file(path, FORMATS[".csv"], registry.actions["csv.import"], warnings.append)
    assert out.rows() == [("1", "x"), ("2", "y")] and "left out" in warnings[0]


@all_formats
def test_text_with_two_escapes_that_share_an_underscore_reads_back_exactly(registry, tmp_path, suffix):
    values = [
        "_x0041_x0042_",
        "a_x0041_x0042_b",
        "poster_x1200_x1600_final",
        "_x0041__x0042_",
        "_x0041_",
        "\x01_x0041_",
        "Serial _x0041\x01 end",
        "_x000D\x1f",
    ]
    frame = text_frame({"a": values, "_x005F_x0041_": values})
    assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


def test_every_escape_in_the_workbook_is_written_as_text(tmp_path):
    path = tmp_path / "a.xlsx"
    write_file(text_frame({"a": ["poster_x1200_x1600_final"]}), path, ".xlsx")
    with zipfile.ZipFile(path) as book:
        assert b"poster_x005F_x1200_x005F_x1600_final" in book.read("xl/worksheets/sheet1.xml")


@all_formats
def test_a_stop_when_the_file_is_complete_leaves_the_old_file(tmp_path, suffix):
    path = tmp_path / f"a{suffix}"
    path.write_bytes(b"old")
    answers = iter([False, True] if suffix == ".xlsx" else [True])
    with pytest.raises(Cancelled):
        write_file(text_frame({"a": ["1"]}), path, suffix, cancelled=lambda: next(answers))
    assert path.read_bytes() == b"old" and [p.name for p in tmp_path.iterdir()] == [path.name]


def test_read_file_in_a_folder_with_a_variable_in_its_name(registry, tmp_path, monkeypatch):
    monkeypatch.setenv("TEMP", r"C:\Temp")
    folder = tmp_path / "proj %TEMP%"
    folder.mkdir()
    path = folder / "a.csv"
    write_file(text_frame({"a": ["1"]}), path, ".csv")
    assert read_file(path, FORMATS[".csv"], registry.actions["csv.import"]).rows() == [("1",)]


# --- the import of other Excel files stays as it was --------------------------------------------


def plain_workbook(path: Path) -> None:
    book = xlsxwriter.Workbook(str(path), {"strings_to_numbers": False})
    sheet = book.add_worksheet()
    for row, values in enumerate([("a", "b"), ("x", "   "), (None, None), ("y", "z")]):
        for column, value in enumerate(values):
            if value is not None:
                sheet.write_string(row, column, value)
    book.close()


@pytest.mark.parametrize(
    "extra",
    [
        {},
        {files.SHEET_KEY: {"file": "other.xlsx", "stamp": "1-2"}},  # another file was chosen in the step
        {files.SHEET_KEY: "yes"},
    ],
)
def test_excel_files_without_the_mark_are_read_as_always(registry, tmp_path, extra):
    path = tmp_path / "plain.xlsx"
    plain_workbook(path)
    out = registry.actions["excel.import"](None, {"source": "file", "file": str(path), **extra}, StubContext())
    assert out.rows() == [("x", None), ("y", "z")]


def test_a_folder_import_is_read_as_always_even_with_an_old_mark(registry, tmp_path):
    plain_workbook(tmp_path / "plain.xlsx")
    params = {
        "source": "folder",
        "folder": str(tmp_path),
        "file_column": False,
        files.SHEET_KEY: {"file": "plain.xlsx"},
    }
    assert registry.actions["excel.import"](None, params, StubContext()).rows() == [("x", None), ("y", "z")]


def test_the_same_workbook_with_the_mark_is_exact(registry, tmp_path):
    path = tmp_path / "plain.xlsx"
    plain_workbook(path)
    out = registry.actions["excel.import"](None, step_params(path, ".xlsx"), StubContext())
    assert out.rows() == [("x", "   "), (None, None), ("y", "z")]


def test_the_versions_of_the_imports_stay(registry):
    assert [registry.actions[name].version for name in ("csv.import", "excel.import", "parquet.import")] == [3, 3, 2]


def test_csv_detection_is_the_same_after_the_move(act, tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    assert act("csv.import", None, {"file": str(path)}).columns == ["a", "b"]
    path.write_text("price\n12,50\n", encoding="utf-8")
    assert act("csv.import", None, {"file": str(path), "detect_types": False}).rows() == [("12,50",)]
    assert files.detect_delimiter("a;b\n1;2\n") == ";"
    assert files.detect_delimiter("skip\na|b\n1|2\n", 1) == "|"
    assert files.detect_delimiter("x\n") == files.ONE_COLUMN
    assert files.detect_delimiter("") == ","


# --- the mark in a project ------------------------------------------------------------------------


@all_formats
def test_a_project_with_a_marked_step_runs_and_keeps_the_mark(registry, tmp_path, suffix):
    frame = text_frame({"Nr": ["00123", None, "1E5"], "Betrag": ["1,50", None, "   "]})
    path = tmp_path / f"data{suffix}"
    write_file(frame, path, suffix)
    project = Project(name="p", path=tmp_path / "p.tafl")
    table = project.add_table("data")
    sheet_format = FORMATS[suffix]
    action = registry.actions[sheet_format.action]
    project.add_step(
        table, action, {**sheet_format.settings, "source": "file", "file": path.name, files.SHEET_KEY: mark(path)}
    )
    project.save()
    assert json.loads((tmp_path / "p.tafl").read_text(encoding="utf-8"))["tables"][0]["steps"][0]["params"][
        files.SHEET_KEY
    ] == mark(path)
    loaded = Project.load(tmp_path / "p.tafl")
    engine = Engine(loaded, registry)
    assert engine.upgrade_steps() == []
    result = engine.evaluate_all(loaded.tables[0].id)[loaded.tables[0].id]
    assert_frame_equal(result.df, frame)
    step = loaded.tables[0].steps[0]
    assert source_format(step.action, step.params) is sheet_format


# --- writing ------------------------------------------------------------------------------------


def test_the_text_files_are_as_a_spreadsheet_program_reads_them(tmp_path):
    frame = text_frame({"a": ["1", "x;y"], "b": ["z\tw", None]})
    write_file(frame, tmp_path / "a.csv", ".csv")
    assert (tmp_path / "a.csv").read_bytes() == '﻿a;b\r\n1;z\tw\r\n"x;y";\r\n'.encode()
    write_file(frame, tmp_path / "a.txt", ".txt")
    assert (tmp_path / "a.txt").read_bytes() == '﻿a\tb\r\n1\t"z\tw"\r\nx;y\t\r\n'.encode()


def test_the_workbook_holds_text_only(tmp_path):
    path = tmp_path / "a.xlsx"
    write_file(text_frame({"Nr": ["00123", "=1+2", "http://x.y/z", "1E5", None]}), path, ".xlsx")
    book = openpyxl.load_workbook(path)
    sheet = book.active
    assert sheet.title == "Sheet1" and sheet.freeze_panes == "A2"
    assert sheet["A1"].font.b
    assert sheet.column_dimensions["A"].width > 10
    assert [cell.value for cell in sheet["A"]] == ["Nr", "00123", "=1+2", "http://x.y/z", "1E5"]
    assert "f" not in {cell.data_type for cell in sheet["A"]}
    assert not sheet["A4"].hyperlink
    book.close()


def test_a_replaced_file_gets_the_new_data_and_a_new_stamp(tmp_path):
    path = tmp_path / "a.csv"
    write_file(text_frame({"a": ["1"]}), path, ".csv")
    stamp = stamp_of(path)
    write_file(text_frame({"a": ["1", "2"]}), path, ".csv")
    assert stamp_of(path) != stamp


def test_the_folder_is_created_and_a_bad_name_is_said_plainly(tmp_path):
    write_file(text_frame({"a": ["1"]}), tmp_path / "new" / "a.csv", ".csv")
    assert (tmp_path / "new" / "a.csv").is_file()
    with pytest.raises(ActionError, match="reserved for a device"):
        write_file(text_frame({"a": ["1"]}), tmp_path / "CON.csv", ".csv")


@pytest.fixture
def own_temp(tmp_path, monkeypatch) -> Path:
    """The temporary folder for this test: what a stop leaves behind there can be seen."""
    folder = tmp_path / "temp"
    folder.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(folder))
    return folder


def left_over(folder: Path) -> list[str]:
    return sorted(path.name for path in folder.iterdir())


def test_a_stop_leaves_the_old_file_and_nothing_else(tmp_path, own_temp):
    target = tmp_path / "out"
    target.mkdir()
    path = target / "a.xlsx"
    path.write_bytes(b"old")
    seen = []

    def stop_soon() -> bool:
        seen.append(1)
        return len(seen) == 3

    with pytest.raises(Cancelled):
        write_file(text_frame({"a": ["x"] * 200_000}), path, ".xlsx", cancelled=stop_soon)
    assert len(seen) == 3
    assert path.read_bytes() == b"old"
    assert left_over(target) == ["a.xlsx"]
    assert left_over(own_temp) == []


def test_a_stop_of_a_new_file_leaves_nothing(tmp_path, own_temp):
    target = tmp_path / "out"
    target.mkdir()
    with pytest.raises(Cancelled):
        write_file(text_frame({"a": ["x"] * 100}), target / "a.xlsx", ".xlsx", cancelled=lambda: True)
    assert left_over(target) == [] and left_over(own_temp) == []


def test_progress_goes_up_to_the_end(tmp_path):
    seen = []
    write_file(
        text_frame({"a": ["x"] * 120_000, "b": ["y"] * 120_000}), tmp_path / "a.xlsx", ".xlsx", progress=seen.append
    )
    assert len(seen) == 6 and seen == sorted(seen) and seen[-1] == 1 and seen[-2] < 1


@pytest.mark.parametrize("failing", ["write", "close"])
def test_an_error_while_writing_a_workbook_leaves_the_old_file(tmp_path, own_temp, monkeypatch, failing):
    target = tmp_path / "out"
    target.mkdir()
    path = target / "a.xlsx"
    path.write_bytes(b"old")

    def full_disk(*args, **kwargs):
        try:
            raise OSError(28, "There is not enough space on the disk")
        except OSError as error:
            if failing == "write":
                raise
            raise xlsxwriter.exceptions.FileCreateError(error) from error  # as xlsxwriter does it

    if failing == "write":
        monkeypatch.setattr(xlsxwriter.worksheet.Worksheet, "write_string", full_disk)
    else:
        monkeypatch.setattr(xlsxwriter.Workbook, "close", full_disk)
    with pytest.raises(ActionError, match=r"^Can't write a.xlsx: There is not enough space on the disk$"):
        write_file(text_frame({"a": ["x"]}), path, ".xlsx")
    assert path.read_bytes() == b"old"
    assert left_over(target) == ["a.xlsx"] and left_over(own_temp) == []


@pytest.mark.parametrize("suffix", [".csv", ".txt", ".parquet"])
def test_an_error_while_writing_leaves_the_old_file(tmp_path, monkeypatch, suffix):
    path = tmp_path / f"a{suffix}"
    path.write_bytes(b"old")

    def fail(*args, **kwargs):
        raise pl.exceptions.ComputeError("the disk is full\nand more")

    monkeypatch.setattr(pl.DataFrame, "write_csv", fail)
    monkeypatch.setattr(pl.DataFrame, "write_parquet", fail)
    with pytest.raises(ActionError) as error:
        write_file(text_frame({"a": ["x"]}), path, suffix)
    assert str(error.value) == f"Can't write a{suffix}: the disk is full"
    assert path.read_bytes() == b"old" and left_over(tmp_path) == [path.name]


@windows
@all_formats
@pytest.mark.parametrize("share_read", [True, False])
def test_a_file_held_open_by_another_program(tmp_path, suffix, share_read):
    path = tmp_path / f"a{suffix}"
    path.write_bytes(b"old")
    with held(path, share_read), pytest.raises(ActionError) as error:
        write_file(text_frame({"a": ["x"]}), path, suffix)
    assert str(error.value) == f"a{suffix} is open in another program (e.g. Excel): close it and run again"
    assert path.read_bytes() == b"old" and left_over(tmp_path) == [path.name]
    write_file(text_frame({"a": ["x"]}), path, suffix)  # the next try works
    assert path.read_bytes() != b"old"


def test_a_read_only_file(tmp_path):
    path = tmp_path / "a.csv"
    path.write_bytes(b"old")
    path.chmod(0o444)
    try:
        with pytest.raises(ActionError, match="a.csv is read-only, so it can't be replaced"):
            write_file(text_frame({"a": ["x"]}), path, ".csv")
    finally:
        path.chmod(0o666)
    assert path.read_bytes() == b"old"


def test_a_warning_reaches_the_caller(tmp_path, monkeypatch):
    path = tmp_path / "a.csv"
    path.write_bytes(b"old")

    def no_new_files(*args, **kwargs):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(files, "_new_temporary", no_new_files)
    warned = []
    write_file(text_frame({"a": ["x"]}), path, ".csv", warn=warned.append)
    assert "written over directly" in warned[0]
    assert b"x" in path.read_bytes()


def test_excel_limits(monkeypatch):
    frame = text_frame({"a": ["x", "y", "z"], "b": ["1", "2", "3"]})
    assert format_problem(frame, ".csv") is None and format_problem(frame, ".parquet") is None
    assert format_problem(frame, ".xlsx") is None
    monkeypatch.setattr(module, "_XLSX_ROWS", 2)
    assert format_problem(frame, ".xlsx") == (
        "An Excel sheet holds at most 2 rows below the column names, and this sheet has 3. Choose CSV, TXT or Parquet."
    )
    monkeypatch.setattr(module, "_XLSX_ROWS", 3)
    monkeypatch.setattr(module, "_XLSX_COLUMNS", 1)
    assert (
        format_problem(frame, ".xlsx")
        == "An Excel sheet holds at most 1 columns, and this sheet has 2. Choose CSV, TXT or Parquet."
    )


def test_cells_too_long_for_excel(tmp_path):
    long = "x" * 32_768
    fits = "x" * 32_767
    frame = text_frame({"a": [fits, None]})
    assert format_problem(frame, ".xlsx") is None
    frame = text_frame({"a": [long, None], "b": [fits, fits]})
    one = "1 cell has more than 32,767 characters, the most an Excel cell holds. Choose CSV, TXT or Parquet."
    assert format_problem(frame, ".xlsx") == one
    frame = text_frame({"a": [long, long], "b": [long, None]})
    assert format_problem(frame, ".xlsx") == (
        "3 cells have more than 32,767 characters, the most an Excel cell holds. Choose CSV, TXT or Parquet."
    )
    assert format_problem(text_frame({long: ["x"]}), ".xlsx") == one
    with pytest.raises(ActionError) as error:  # nothing cuts a text without saying so
        write_file(frame, tmp_path / "a.xlsx", ".xlsx")
    assert str(error.value).startswith("3 cells have")
    assert not (tmp_path / "a.xlsx").exists()


def test_the_long_cell_reads_back_in_the_other_formats(registry, tmp_path):
    frame = text_frame({"a": ["x" * 40_000, None, "y"]})
    for suffix in (".csv", ".txt", ".parquet"):
        assert_frame_equal(read_back(registry, frame, tmp_path, suffix), frame)


def test_the_zip_extension_is_used_only_when_the_file_needs_it(tmp_path, monkeypatch):
    calls = []
    original = xlsxwriter.Workbook.use_zip64
    monkeypatch.setattr(xlsxwriter.Workbook, "use_zip64", lambda book: (calls.append(1), original(book)))
    frame = text_frame({"a": ["x", None]})
    write_file(frame, tmp_path / "a.xlsx", ".xlsx")
    assert not calls
    monkeypatch.setattr(module, "_XLSX_BIG", 10)
    write_file(frame, tmp_path / "a.xlsx", ".xlsx")
    assert calls == [1]


def test_the_size_of_the_xml_counts_the_tags():
    frame = text_frame({"a": ["x"] * 10, "b": [None] * 10})
    assert not module._big_xml(frame)
    many = pl.DataFrame({f"c{i}": pl.repeat("x", 5_000_000, dtype=pl.String, eager=True) for i in range(8)})
    assert many.estimated_size() < 2 << 30  # the size of the data says nothing about the XML
    assert module._big_xml(many)
