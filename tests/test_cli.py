import hashlib
import importlib.util
import json
import os
import re
import signal
import subprocess
import sys
import threading
import time
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from tafl import cli
from tafl.api import expressions
from tafl.cli import Console, main
from tafl.core import Parameter, Project
from tafl.core.cancel import CancelEvent

EXAMPLES = Path(__file__).resolve().parents[1] / "examples" / "make_examples.py"
TAFL_EXE = Path(__file__).resolve().parents[1] / "packaging" / "tafl_cli.py"  # how tafl.exe starts


@pytest.fixture
def demo(tmp_path) -> Path:
    """The demo project from examples/, built in a temporary folder."""
    spec = importlib.util.spec_from_file_location("make_examples", EXAMPLES)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    (tmp_path / "data").mkdir()
    module.write_customers(tmp_path / "data" / "customers.xlsx")
    module.write_orders(tmp_path / "data" / "orders.csv")
    return module.build_project().save(tmp_path / "demo.tafl")


def test_run_demo_project(demo, capsys):
    assert main(["run", str(demo)]) == 0
    output = capsys.readouterr().out
    assert "Done in" in output and "Revenue by customer" in output

    revenue = openpyxl.load_workbook(demo.parent / "output" / "revenue.xlsx")["Revenue"]
    header = [c.value for c in revenue[1]]
    assert header == ["Name", "City", "Revenue", "Orders", "Last order", "Share"]
    assert revenue["C2"].number_format == "#,##0.00 €"
    assert revenue["F2"].number_format == "0.0%"
    shares = [row[5].value for row in revenue.iter_rows(min_row=2)]
    assert sum(shares) == pytest.approx(1.0)

    orders_csv = (demo.parent / "output" / "orders.csv").read_text(encoding="utf-8-sig").splitlines()
    assert orders_csv[0] == "OrderNo;CustomerNo;Date;Net;Status;Name;City;Gross;Month;Size;Order of customer"


def test_preview_run_writes_nothing(demo, capsys):
    assert main(["run", str(demo), "--preview"]) == 0
    assert "(runs only when the project is run)" in capsys.readouterr().out
    assert not (demo.parent / "output").exists()


def test_show_a_step(demo, capsys):
    assert main(["show", str(demo), "--table", "Customers", "--step", "3"]) == 0
    output = capsys.readouterr().out
    assert "6 rows × 3 columns" in output
    assert "Bäckerei Weiß" in output  # repeated spaces collapsed
    assert main(["show", str(demo), "--table", "Customers", "--step", "9"]) == 2


def test_failing_project_returns_1(demo, capsys):
    text = demo.read_text(encoding="utf-8").replace("ROUND([Net] * 1.19, 2)", "ROUND([Nett] * 1.19, 2)")
    demo.write_text(text, encoding="utf-8")
    assert main(["run", str(demo)]) == 1
    output = capsys.readouterr().out
    assert "column [Nett] not found (did you mean [Net]?)" in output
    assert "Table 'Orders' has an error" in output  # Revenue by customer depends on Orders


def test_plugins_and_missing_project(capsys, tmp_path):
    assert main(["plugins"]) == 0
    assert "calc.enumerate" in capsys.readouterr().out
    assert main(["run", str(tmp_path / "missing.tafl")]) == 2


def test_a_console_without_unicode_gets_words_for_the_symbols(capsys):
    console = Console()
    console.unicode = False  # e.g. output to a file in the Windows code page
    console.print("4. Sort by Revenue ↓, Name → 2 × 3")
    assert capsys.readouterr().out == "4. Sort by Revenue desc, Name -> 2 x 3\n"


def test_long_step_summaries_are_shortened(capsys, monkeypatch):
    console, flushed = Console(), []
    monkeypatch.setattr("builtins.print", lambda *args, **kwargs: flushed.append(kwargs.get("flush")))
    console.step(3, "Keep rows where Kunde is one of " + ", ".join(f"K{i:05d}" for i in range(20_000)), None)
    console.unicode = False
    console.step(4, "Sort by Name", None)
    monkeypatch.undo()
    assert flushed == [True, True]  # each line reaches a log file at once, also if the run is ended then
    console.unicode = True
    console.step(3, "Keep rows where Kunde is one of " + ", ".join(f"K{i:05d}" for i in range(20_000)), None)
    line = capsys.readouterr().out.rstrip("\n")
    assert line.startswith("        3. Keep rows where Kunde is one of K00000, ") and line.endswith("…")
    assert len(line) == len("        3. ") + 300


def test_without_unicode_the_columns_of_a_table_line_up(capsys):
    console = Console()
    console.unicode = False
    console.table(pl.DataFrame({"From → to": ["A → B", "long…", "x"], "n": [1, 2, 3]}), 2)
    lines = capsys.readouterr().out.splitlines()
    assert len({len(line) for line in lines}) == 1, lines
    assert "From -> to" in lines[1] and "| A -> B " in lines[5] and "| long... " in lines[6]


@pytest.fixture
def plugins(tmp_path, monkeypatch) -> Path:
    """A plugin folder with actions that fail in different ways; no plugins of this computer."""
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    monkeypatch.delenv("TAFL_PLUGINS", raising=False)
    folder = tmp_path / "plugins"
    folder.mkdir()
    (folder / "quits.py").write_text("import sys\nsys.exit('plugin says bye')\n", encoding="utf-8")
    (folder / "failing.py").write_text(FAILING, encoding="utf-8")
    return folder


FAILING = '''
import signal
import threading
import time

import polars as pl

from tafl.api import ActionError, action, expressions


@action("demo.interrupt", "Interrupt")
def interrupt(df, p, ctx):
    raise KeyboardInterrupt


@action("demo.ctrl_c", "Ctrl+C")
def ctrl_c(df, p, ctx):
    signal.raise_signal(signal.SIGINT)
    time.sleep(1)  # the handler runs before this ends
    return df


@action("demo.slow_failure", "Slow failure")
def slow_failure(df, p, ctx):
    time.sleep(0.2)
    raise ActionError("The server did not answer")


@action("demo.slow_formula", "Slow formula")
def slow_formula(df, p, ctx):
    threading.Timer(0.3, signal.raise_signal, (signal.SIGINT,)).start()  # Ctrl+C while it runs
    return expressions.filter_rows(df, "(SELECT COUNT(*) FROM range(3000000000) t(i) WHERE i % 7 = 8) = 0")


def _interrupt(*_):
    raise KeyboardInterrupt


@action("demo.old", "Old", version=2, migrate=_interrupt)
def old(df, p, ctx):
    return df


@action("demo.titled", "Titled", summary=_interrupt)
def titled(df, p, ctx):
    return df


@action("demo.source", "Source", kind="source", fingerprint=_interrupt)
def source(df, p, ctx):
    return pl.DataFrame({"a": [1]})
'''


def _plugin_project(folder: Path, *actions: str) -> Path:
    (folder / "orders.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="Plugins")
    table = project.add_table("Orders")
    project.add_step(table, "csv.import", {"file": "orders.csv"})
    for action_id in actions:
        project.add_step(table, action_id, {})
    return project.save(folder / "plugins.tafl")


def test_plugins_that_try_to_end_the_program(plugins, tmp_path, capsys):
    path = _plugin_project(tmp_path, "demo.interrupt")
    assert main(["--plugins", str(plugins), "run", str(path)]) == 1
    output = capsys.readouterr()
    assert "Warning: Plugin 'quits' failed to load: SystemExit: plugin says bye" in output.err
    assert "The plugin 'failing' tried to end the program (KeyboardInterrupt)" in output.out
    assert "Failed in" in output.out


def test_ctrl_c_stops_a_run(plugins, tmp_path, capsys):
    before = signal.getsignal(signal.SIGINT)
    path = _plugin_project(tmp_path, "demo.ctrl_c", "demo.interrupt")
    assert main(["--plugins", str(plugins), "run", str(path)]) == 130
    output = capsys.readouterr()
    assert "Error: Stopped" in output.err and "Failed in" not in output.out
    assert signal.getsignal(signal.SIGINT) is before


def test_after_ctrl_c_what_ran_is_printed(plugins, tmp_path, capsys):
    """The tables that were done (with their exports), and the steps of the stopped table."""
    path = _plugin_project(tmp_path, "csv.export", "demo.ctrl_c", "demo.interrupt")
    project = Project.load(path)
    project.tables[0].steps[1].params["file"] = "out/orders.csv"
    first = project.add_table("First")
    project.tables.insert(0, project.tables.pop())  # calculated first
    project.add_step(first, "csv.import", {"file": "orders.csv"})
    project.add_step(first, "csv.export", {"file": "out/first.csv"})
    project.save(path)
    assert main(["--plugins", str(plugins), "run", str(path)]) == 130
    output = capsys.readouterr()
    lines = output.out.splitlines()
    assert lines[0] == "First" and lines[3].startswith("  ✓ 2. Export to first.csv   1 row × 1 column")
    assert lines[4] == f"          i Wrote 1 row to {tmp_path / 'out' / 'first.csv'}"
    assert lines[5] == "Orders" and lines[8].startswith("  ✓ 2. Export to orders.csv   1 row × 1 column")
    assert lines[10:12] == ["  · 3. Ctrl+C   (stopped)", "        4. Interrupt"]
    assert re.fullmatch(r"Stopped after \d+\.\d\ds", lines[12]) and len(lines) == 13
    assert output.err.endswith("\nError: Stopped\n")
    assert (tmp_path / "out" / "first.csv").exists() and (tmp_path / "out" / "orders.csv").exists()


def test_ctrl_c_while_a_formula_runs_stops_the_run(plugins, tmp_path, capsys):
    """DuckDB reports the interrupted query as an error: that is the stop, not a step error."""
    path = _plugin_project(tmp_path, "demo.slow_formula", "csv.export")
    before = signal.signal(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    try:
        started = time.perf_counter()
        assert main(["--plugins", str(plugins), "run", str(path)]) == 130
        assert time.perf_counter() - started < 10
    finally:
        signal.signal(signal.SIGINT, before)
    output = capsys.readouterr()
    assert "Query interrupted" not in output.out and "Failed in" not in output.out
    assert "  · 2. Slow formula   (stopped)" in output.out and output.err.endswith("Error: Stopped\n")


def test_ctrl_c_is_seen_at_once_while_python_waits_for_a_library():
    """A library call that doesn't look for signals (as a DuckDB query sometimes does) holds back
    the interruption until it is over, but the stop is set at once by another thread."""
    stop, seen = CancelEvent(), []
    stop.when_set(lambda: seen.append((threading.current_thread().name, time.perf_counter())))
    before = signal.signal(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    started = time.perf_counter()
    try:
        with pytest.raises(KeyboardInterrupt), cli._ctrl_c(stop):
            threading.Timer(0.05, signal.raise_signal, (signal.SIGINT,)).start()
            hashlib.pbkdf2_hmac("sha256", b"x", b"y", 2_000_000)  # a second or so without looking for signals
    finally:
        signal.signal(signal.SIGINT, before)
    ended = time.perf_counter()  # the interruption came when the call was over
    assert seen and seen[0][0] == "ctrl-c" and seen[0][1] - started < (ended - started) / 2
    assert signal.set_wakeup_fd(-1) == -1  # the watching ended


def test_ctrl_c_while_the_stop_event_holds_its_lock_doesnt_hang():
    """The Ctrl+C handler runs in the main thread, also in the middle of the stop event's own
    code: it must not wait for the event's lock, which that code holds (it hung there)."""
    stop, seen = CancelEvent(), []
    stop.when_set(lambda: seen.append(threading.current_thread().name))
    before = signal.signal(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    try:
        with pytest.raises(KeyboardInterrupt), cli._ctrl_c(stop):
            with stop._lock:  # as when the handler comes while a query's stopping is removed
                signal.raise_signal(signal.SIGINT)
                time.sleep(0.2)  # the handler runs here: it must return at once
    finally:
        signal.signal(signal.SIGINT, before)
    assert stop.is_set()
    deadline = time.perf_counter() + 5
    while not seen and time.perf_counter() < deadline:  # the watching thread calls the callbacks
        time.sleep(0.01)
    assert seen == ["ctrl-c"]


def test_a_stop_event_set_from_a_signal():
    stop, seen = CancelEvent(), []
    stop.when_set(lambda: seen.append("first"))
    with stop._lock:
        stop.set_from_signal()  # takes no lock
    assert stop.is_set() and seen == []
    stop.when_set(lambda: seen.append("later"))  # set already: at once
    stop.set()  # the rest: the callbacks
    assert seen == ["later", "first"]
    stop.clear()
    assert not stop.is_set()


def _one_table(folder: Path) -> Path:
    """A project that imports orders.csv and writes out.csv."""
    (folder / "orders.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="Orders")
    table = project.add_table("Orders")
    project.add_step(table, "csv.import", {"file": "orders.csv"})
    project.add_step(table, "csv.export", {"file": "out.csv"})
    return project.save(folder / "orders.tafl")


def test_as_the_program_ctrl_c_is_ignored_once_the_work_is_over(tmp_path, monkeypatch, capsys):
    """tafl.exe, tafl and python -m tafl: a Ctrl+C while Python ends can make the program hang,
    so from the end of the work on Ctrl+C does nothing."""
    monkeypatch.setattr(sys, "argv", ["tafl", "run", str(_one_table(tmp_path))])
    before = signal.signal(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    try:
        assert main() == 0
        assert signal.getsignal(signal.SIGINT) is signal.SIG_IGN
        signal.raise_signal(signal.SIGINT)  # does nothing
    finally:
        signal.signal(signal.SIGINT, before)
    assert "Done in" in capsys.readouterr().out
    assert main(["run", str(tmp_path / "orders.tafl")]) == 0  # called with a command line: as before
    assert signal.getsignal(signal.SIGINT) is before


def test_a_ctrl_c_as_the_program_ends_changes_nothing(tmp_path, monkeypatch, capsys):
    """Ctrl+C just after "Done in" (e.g. held down): the result stays, and Ctrl+C is still
    switched off, also when one comes just as the handler is changed."""
    monkeypatch.setattr(sys, "argv", ["tafl", "run", str(_one_table(tmp_path))])
    switch = signal.signal

    def switching(signum, handler):
        if handler is signal.SIG_IGN:
            signal.raise_signal(signal.SIGINT)  # Ctrl+C just then
        return switch(signum, handler)

    before = switch(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    try:
        monkeypatch.setattr(signal, "signal", switching)
        assert main() == 0
        monkeypatch.setattr(signal, "signal", switch)
        assert signal.getsignal(signal.SIGINT) is signal.SIG_IGN
    finally:
        switch(signal.SIGINT, before)
    output = capsys.readouterr()
    assert "Done in" in output.out and "Stopped" not in output.err


def test_an_interruption_once_the_work_is_over_keeps_the_result(tmp_path, monkeypatch, capsys):
    """A KeyboardInterrupt that comes as the command ends (Ctrl+C just after "Done in")."""
    watch = cli._watch_for_ctrl_c

    def watching(stop):
        end = watch(stop)

        def ending():
            end()
            raise KeyboardInterrupt

        return ending

    monkeypatch.setattr(cli, "_watch_for_ctrl_c", watching)
    before = signal.getsignal(signal.SIGINT)
    assert main(["run", str(_one_table(tmp_path))]) == 0
    output = capsys.readouterr()
    assert "Done in" in output.out and "Stopped" not in output.err
    assert signal.getsignal(signal.SIGINT) is before


def test_a_ctrl_c_before_the_command_stops_it_at_once(tmp_path, capsys):
    """E.g. while tafl.exe loads its modules: packaging/tafl_cli.py keeps it for main()."""
    path, early = _one_table(tmp_path), threading.Event()
    before = signal.getsignal(signal.SIGINT)
    assert main(["run", str(path)], ctrl_c=early) == 0  # not set: the command runs
    assert "Done in" in capsys.readouterr().out
    (tmp_path / "out.csv").unlink()
    early.set()
    assert main(["run", str(path)], ctrl_c=early) == 130
    assert capsys.readouterr() == ("", "Error: Stopped\n")
    assert not (tmp_path / "out.csv").exists()
    assert signal.getsignal(signal.SIGINT) is before


START = '''
import builtins
import multiprocessing
import runpy
import signal
import sys

script, how, *args = sys.argv[1:]
load = builtins.__import__


def loading(name, *more, **options):
    if name == "tafl.cli":  # Ctrl+C while the program loads its modules
        builtins.__import__ = load
        signal.raise_signal(signal.SIGINT)
    return load(name, *more, **options)


def worker():
    print("worker", signal.getsignal(signal.SIGINT) is signal.default_int_handler)
    sys.exit(0)


if how == "ctrl-c":
    builtins.__import__ = loading
if how == "worker":
    multiprocessing.freeze_support = worker
    args = ["--multiprocessing-fork", "parent_pid=1", "pipe_handle=2"]
sys.argv = ["tafl", *args]
runpy.run_path(script, run_name="__main__")
'''


def test_tafl_exe_keeps_a_ctrl_c_that_comes_while_it_starts(tmp_path):
    """Without it, Ctrl+C in the first moment gives a KeyboardInterrupt in the middle of loading
    a library: a traceback, or a crash."""
    path = _one_table(tmp_path)
    (tmp_path / "start.py").write_text(START, encoding="utf-8")
    environment = {**os.environ, "APPDATA": str(tmp_path / "appdata")}  # no plugins of this computer
    environment.pop("TAFL_PLUGINS", None)

    def start(how: str, *args: str) -> subprocess.CompletedProcess:
        command = [sys.executable, "-I", str(tmp_path / "start.py"), str(TAFL_EXE), how, *args]
        return subprocess.run(command, capture_output=True, text=True, errors="replace", env=environment, timeout=120)

    stopped = start("ctrl-c", "run", str(path))
    assert (stopped.returncode, stopped.stderr) == (130, "Error: Stopped\n")
    assert not (tmp_path / "out.csv").exists()
    done = start("plain", "run", str(path))
    assert done.returncode == 0 and "Done in" in done.stdout and (tmp_path / "out.csv").exists()
    worker = start("worker")  # a plugin's worker process: Python's own handling of Ctrl+C
    assert (worker.returncode, worker.stdout) == (0, "worker True\n")


SLOW = "=(SELECT COUNT(*) FROM range(30000000000) t(i) WHERE i % 7 = 8)"  # takes minutes


@pytest.mark.parametrize("ctrl_c", ["signal", "watching"])
def test_ctrl_c_during_a_parameter_formula_stops_the_run(ctrl_c, tmp_path, monkeypatch, capsys):
    """At once, and as the stop: no error of the parameter. Also when only the thread that
    watches for Ctrl+C sets the stop (Python's own handling may not come while DuckDB runs)."""
    (tmp_path / "orders.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="Slow", parameters=[Parameter("Slow", SLOW)])
    table = project.add_table("Orders")
    project.add_step(table, "csv.import", {"file": "orders.csv"})
    project.add_step(table, "calc.columns", {"columns": [{"name": "X", "formula": "{Slow}"}]})
    project.add_step(table, "csv.export", {"file": "out.csv"})
    path = project.save(tmp_path / "slow.tafl")

    running, stops, pressed, evaluate = threading.Event(), [], [], expressions.evaluate

    def evaluating(*args, **kwargs):  # the first formula is the parameter's
        running.set()
        return evaluate(*args, **kwargs)

    class Stop(CancelEvent):
        def __init__(self) -> None:
            super().__init__()
            stops.append(self)

    def press() -> None:
        running.wait(30)
        time.sleep(0.3)
        pressed.append(time.perf_counter())
        if ctrl_c == "signal":
            signal.raise_signal(signal.SIGINT)
        else:
            stops[0].set()

    monkeypatch.setattr(expressions, "evaluate", evaluating)
    monkeypatch.setattr(cli, "CancelEvent", Stop)
    before = signal.signal(signal.SIGINT, lambda *_: None)  # never a Ctrl+C for pytest itself
    try:
        threading.Thread(target=press, daemon=True).start()
        assert main(["run", str(path)]) == 130
        assert time.perf_counter() - pressed[0] < 3  # not the 5 seconds a parameter's formula may take
    finally:
        signal.signal(signal.SIGINT, before)
    output = capsys.readouterr()
    assert re.fullmatch(r"Stopped after \d+\.\d\ds\n", output.out) and output.err == "Error: Stopped\n"
    assert not (tmp_path / "out.csv").exists()


def test_keyboard_interrupt_from_plugin_functions_is_their_error(plugins, tmp_path, capsys):
    """A migrate, summary or fingerprint function that raises KeyboardInterrupt doesn't stop the run."""
    path = _plugin_project(tmp_path, "demo.titled", "demo.old")
    project = Project.load(path)
    project.add_step(project.add_table("Other"), "demo.source", {})
    project.save(path)
    assert main(["--plugins", str(plugins), "run", str(path)]) == 1
    output = capsys.readouterr()
    assert "  ✓ 2. Titled   1 row × 1 column" in output.out
    assert "  ✗ 3. Old\n" in output.out
    assert "Could not read the step settings: The plugin 'failing' tried to end the program (KeyboardInterrupt)" in output.out
    assert "  ✓ 1. Source   1 row × 1 column" in output.out
    assert "Failed in" in output.out and "Stopped" not in output.err


def test_a_failed_step_shows_how_long_it_ran(plugins, tmp_path, capsys):
    path = _plugin_project(tmp_path, "demo.slow_failure")
    assert main(["--plugins", str(plugins), "run", str(path)]) == 1
    assert re.search(r"✗ 2\. Slow failure   0\.[2-9]\ds\n +The server did not answer", capsys.readouterr().out)


def test_plain_messages_for_a_project_that_isnt_there(tmp_path, capsys):
    assert main(["run", str(tmp_path / "missing.tafl")]) == 2
    assert capsys.readouterr().err == f"Error: There is no project file {tmp_path / 'missing.tafl'}\n"
    assert main(["show", str(tmp_path), "--table", "Orders"]) == 2
    assert capsys.readouterr().err == f"Error: {tmp_path} is a folder, not a project file\n"


def test_plugin_folders_that_dont_exist_get_a_warning(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    monkeypatch.delenv("TAFL_PLUGINS", raising=False)
    assert main(["--plugins", str(tmp_path / "missing"), "plugins"]) == 0
    assert f"Warning: The plugin folder {tmp_path / 'missing'} does not exist" in capsys.readouterr().err


@pytest.mark.parametrize("rows", ["0", "-1", "x"])
def test_rows_needs_at_least_one(demo, rows, capsys):
    with pytest.raises(SystemExit) as stopped:
        main(["show", str(demo), "--table", "Customers", "--rows", rows])
    assert stopped.value.code == 2
    assert "argument --rows: " in capsys.readouterr().err


def test_the_help_names_the_exit_codes(capsys):
    for command in ([], ["run"], ["show"], ["plugins"]):
        with pytest.raises(SystemExit):
            main([*command, "--help"])
        output = capsys.readouterr().out
        assert "exit codes:\n  0    " in output
    with pytest.raises(SystemExit):
        main(["show", "--help"])
    output = capsys.readouterr().out
    assert "project           the project file (*.tafl)" in output and "the table to show" in output
    assert "  130  stopped with Ctrl+C" in output


def test_tables_that_are_off_are_skipped(demo, capsys):
    data = json.loads(demo.read_text(encoding="utf-8"))
    data["tables"][0]["enabled"] = False  # Customers: Orders uses it, and Revenue uses Orders
    demo.write_text(json.dumps(data), encoding="utf-8")
    assert main(["run", str(demo)]) == 0  # skipped tables are not errors
    output = capsys.readouterr().out
    assert "skipped: turned off" in output
    assert "skipped: uses “Customers”, which is turned off" in output
    assert "skipped: uses “Orders”, which is skipped" in output
    assert not (demo.parent / "output").exists()

    assert main(["show", str(demo), "--table", "Orders"]) == 0
    output = capsys.readouterr().out
    assert "Orders\n  ○ skipped: uses “Customers”, which is turned off\n        1. Import orders.csv\n" in output


def test_one_row(tmp_path, capsys):
    (tmp_path / "one.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="One")
    project.add_step(project.add_table("One"), "csv.import", {"file": "one.csv"})
    path = project.save(tmp_path / "one.tafl")
    assert main(["show", str(path), "--table", "One"]) == 0
    output = capsys.readouterr().out
    assert "1. Import one.csv   1 row × 1 column   " in output and "\n1 row × 1 column\n" in output


def test_a_name_that_cant_be_used_but_nothing_uses_only_gets_a_warning(tmp_path, capsys):
    (tmp_path / "one.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="Names", parameters=[Parameter("Größe", "5")])
    step = project.add_step(project.add_table("One"), "csv.import", {"file": "one.csv"})
    path = project.save(tmp_path / "names.tafl")
    assert main(["run", str(path)]) == 0
    assert "Warning: The parameter name 'Größe' can't be used: only letters A-Z" in capsys.readouterr().err

    step.params["file"] = "{Größe}.csv"
    project.save()
    assert main(["run", str(path)]) == 2
    assert "Error: The parameter name 'Größe' can't be used: only letters A-Z" in capsys.readouterr().err


def test_a_formula_that_fails_with_a_set_value_names_it(tmp_path, capsys):
    project = Project(name="Months", parameters=[Parameter("Day", "=TODAY() - 1"), Parameter("Month", "=STRFTIME({Day}, '%Y-%m')")])
    (tmp_path / "one.csv").write_text("Id\n1\n", encoding="utf-8")
    project.add_step(project.add_table("Months"), "csv.import", {"file": "one.csv"})
    path = project.save(tmp_path / "months.tafl")
    main(["run", str(path), "--preview", "--set", "Day=yesterday"])
    assert "({Day} was given as 'yesterday', a text, but the project's value is a date)" in capsys.readouterr().out


def test_parameters_and_set(tmp_path, capsys):
    (tmp_path / "orders_2026-10-01.csv").write_text("Id\n1\n2\n", encoding="utf-8")
    (tmp_path / "orders_2026-10-02.csv").write_text("Id\n1\n2\n3\n", encoding="utf-8")
    project = Project(name="Daily", parameters=[Parameter("Day", "2026-10-01"), Parameter("Next", "={Day} + 1")])
    table = project.add_table("Orders")
    project.add_step(table, "csv.import", {"file": "orders_{Day}.csv"})
    project.add_step(table, "csv.export", {"file": "out_{Next}.csv"})
    path = project.save(tmp_path / "daily.tafl")

    assert main(["run", str(path)]) == 0
    output = capsys.readouterr().out
    assert "Parameters\n  {Day} = 2026-10-01\n  {Next} = 2026-10-02\n" in output
    assert (tmp_path / "out_2026-10-02.csv").read_text(encoding="utf-8-sig").splitlines() == ["Id", "1", "2"]

    assert main(["run", str(path), "--set", "day=2026-10-02"]) == 0
    assert "{Day} = 2026-10-02" in capsys.readouterr().out
    assert (tmp_path / "out_2026-10-03.csv").read_text(encoding="utf-8-sig").splitlines() == ["Id", "1", "2", "3"]

    assert main(["show", str(path), "--table", "Orders", "--set", "Day=2026-10-02", "--set", "Next==1 +"]) == 1
    output = capsys.readouterr().out
    assert "{Next} = error: the formula ends too early" in output
    assert "The parameter {Next} has an error" in output  # the export uses it: an error even in a preview

    assert main(["run", str(path), "--set", "Month=10"]) == 2
    assert "The project has no parameter 'Month' (its parameters: Day, Next)" in capsys.readouterr().err
    assert main(["run", str(path), "--set", "Day"]) == 2
    assert "--set needs NAME=VALUE, not 'Day'" in capsys.readouterr().err
    assert main(["run", str(path), "--set", "Day=2026-09-30"]) == 1  # no such file


def test_the_listed_values_are_those_the_run_used(tmp_path, capsys):
    (tmp_path / "orders.csv").write_text("Id\n1\n", encoding="utf-8")
    project = Project(name="Random", parameters=[Parameter("Version", "1.0"), Parameter("R", "=CAST(RANDOM() * 1e9 AS BIGINT)")])
    table = project.add_table("Orders")
    project.add_step(table, "csv.import", {"file": "orders.csv"})
    project.add_step(table, "csv.export", {"file": "out_v{Version}_{R}.csv"})
    path = project.save(tmp_path / "random.tafl")

    assert main(["run", str(path)]) == 0
    output = capsys.readouterr().out
    listed = re.search(r"Parameters\n  \{Version\} = 1\.0\n  \{R\} = (\d+)\n", output)
    assert listed and output.index("Parameters") < output.index("Orders")
    assert [p.name for p in tmp_path.glob("out_*.csv")] == [f"out_v1.0_{listed[1]}.csv"]

    assert main(["show", str(path), "--table", "Orders"]) == 0
    output = capsys.readouterr().out
    assert output.index("Parameters") < output.index("Orders") < output.index(" row × ")
